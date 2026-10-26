"""Resume parser: PDF and DOCX extraction."""
from __future__ import annotations

import io
import logging
import os
import re
from typing import Any, Dict

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Regex helpers for structured extraction
# ---------------------------------------------------------------------------

_EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
_PHONE_RE = re.compile(r"(\+?\d[\d\s\-().]{7,}\d)")
_URL_RE = re.compile(r"https?://[^\s]+")

_SECTION_HEADERS = {
    "summary": re.compile(r"^(summary|objective|profile|about me)\s*$", re.IGNORECASE | re.MULTILINE),
    "skills": re.compile(r"^(skills|technical skills|core competencies|technologies)\s*$", re.IGNORECASE | re.MULTILINE),
    "experience": re.compile(r"^(experience|work experience|employment|professional experience)\s*$", re.IGNORECASE | re.MULTILINE),
    "education": re.compile(r"^(education|academic background|qualifications)\s*$", re.IGNORECASE | re.MULTILINE),
    "projects": re.compile(r"^(projects|personal projects|academic projects)\s*$", re.IGNORECASE | re.MULTILINE),
}


def _extract_contact_info(text: str) -> Dict[str, str]:
    """Extract contact details from raw text."""
    emails = _EMAIL_RE.findall(text)
    phones = _PHONE_RE.findall(text)
    urls = _URL_RE.findall(text)
    return {
        "email": emails[0] if emails else "",
        "phone": phones[0] if phones else "",
        "urls": urls[:5],
    }


def _extract_sections(text: str) -> Dict[str, str]:
    """Split text into named sections based on section headers."""
    sections: Dict[str, str] = {}
    lines = text.split("\n")
    current_section = "header"
    section_text: list[str] = []

    for line in lines:
        matched = False
        for section_name, pattern in _SECTION_HEADERS.items():
            if pattern.match(line.strip()):
                sections[current_section] = "\n".join(section_text).strip()
                current_section = section_name
                section_text = []
                matched = True
                break
        if not matched:
            section_text.append(line)

    sections[current_section] = "\n".join(section_text).strip()
    return sections


COMMON_TECH_CATALOG = [
    "Python", "JavaScript", "TypeScript", "React", "Next.js", "Node.js", "HTML", "CSS", "SQL",
    "PostgreSQL", "MongoDB", "MySQL", "Git", "GitHub", "Docker", "Kubernetes", "AWS", "Azure",
    "GCP", "Linux", "FastAPI", "Django", "Flask", "Tailwind", "Java", "C++", "C", "C#",
    ".NET", "PHP", "Ruby", "Rust", "Go", "Kotlin", "Swift", "PyTorch", "TensorFlow", "Keras",
    "Scikit-Learn", "Pandas", "NumPy", "Matplotlib", "Figma", "UI/UX", "GraphQL", "REST API",
    "Redux", "Express", "Vite", "OpenCV", "NLP", "LLM", "Prompt Engineering", "Machine Learning",
    "Deep Learning", "Data Analysis", "Artificial Intelligence", "Tableau", "Power BI", "Excel",
]


def _extract_name_from_header(header_text: str, fallback_filename: str = "") -> str:
    """Heuristic: look for candidate name in header lines or top text."""
    for line in header_text.split("\n"):
        clean = line.strip()
        if (
            clean
            and 1 <= len(clean.split()) <= 4
            and len(clean) < 40
            and not _EMAIL_RE.search(clean)
            and not re.search(r"\d|http|github|linkedin|resume|cv", clean, re.IGNORECASE)
        ):
            return clean

    # Clean filename fallback if provided
    if fallback_filename:
        base = re.sub(r"\.(pdf|docx)$", "", fallback_filename, flags=re.IGNORECASE)
        clean_name = re.sub(r"(professional|master|updated|latest|resume|cv)", "", base, flags=re.IGNORECASE)
        clean_name = re.sub(r"[_\-]+", " ", clean_name).strip()
        if clean_name and len(clean_name) >= 3:
            return clean_name.title()

    return ""


def _build_parsed_data(text: str, filename: str = "") -> Dict[str, Any]:
    """Build structured parsed_data dict from raw text."""
    contact = _extract_contact_info(text)
    sections = _extract_sections(text)

    header_text = sections.get("header", "") or "\n".join(text.split("\n")[:10])
    name = _extract_name_from_header(header_text, filename)

    # Skills: split by commas, bullets, or newlines
    skills_text = sections.get("skills", "")
    skill_tokens = re.split(r"[,•|\n\t]+", skills_text)
    extracted_skills = [s.strip() for s in skill_tokens if 1 < len(s.strip()) < 50]

    # Keyword scan across the whole document for common tech skills
    text_lower = text.lower()
    for tech in COMMON_TECH_CATALOG:
        pattern = r"\b" + re.escape(tech.lower()) + r"\b"
        if re.search(pattern, text_lower):
            if not any(tech.lower() == s.lower() for s in extracted_skills):
                extracted_skills.append(tech)

    return {
        "name": name,
        "email": contact["email"],
        "phone": contact["phone"],
        "urls": contact["urls"],
        "skills": extracted_skills,
        "summary": sections.get("summary", ""),
        "experience": sections.get("experience", ""),
        "education": sections.get("education", ""),
        "projects": sections.get("projects", ""),
        "raw_sections": {k: v for k, v in sections.items() if k not in ("skills", "summary", "experience", "education", "projects")},
    }


# ---------------------------------------------------------------------------
# PDF parser (pdfminer.six)
# ---------------------------------------------------------------------------

def parse_pdf(file_path: str) -> Dict[str, Any]:
    """Extract text from a PDF file and return structured parsed data.

    Args:
        file_path: Absolute path to the PDF file.

    Returns:
        dict with keys: raw_text, name, email, phone, skills, summary,
        experience, education, projects, urls.
    """
    try:
        from pdfminer.high_level import extract_text as pdf_extract_text
        raw_text = pdf_extract_text(file_path)
    except Exception as exc:
        logger.warning("pdfminer extraction failed for %s: %s", file_path, exc)
        raw_text = ""

    if not raw_text.strip():
        return {"raw_text": "", "error": "Could not extract text from PDF."}

    filename = os.path.basename(file_path)
    parsed = _build_parsed_data(raw_text, filename=filename)
    parsed["raw_text"] = raw_text
    return parsed


# ---------------------------------------------------------------------------
# DOCX parser (python-docx)
# ---------------------------------------------------------------------------

def parse_docx(file_path: str) -> Dict[str, Any]:
    """Extract text from a DOCX file and return structured parsed data.

    Args:
        file_path: Absolute path to the DOCX file.

    Returns:
        dict with keys: raw_text, name, email, phone, skills, summary,
        experience, education, projects, urls.
    """
    try:
        import docx
        doc = docx.Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        raw_text = "\n".join(paragraphs)
    except Exception as exc:
        logger.warning("python-docx extraction failed for %s: %s", file_path, exc)
        raw_text = ""

    if not raw_text.strip():
        return {"raw_text": "", "error": "Could not extract text from DOCX."}

    filename = os.path.basename(file_path)
    parsed = _build_parsed_data(raw_text, filename=filename)
    parsed["raw_text"] = raw_text
    return parsed
